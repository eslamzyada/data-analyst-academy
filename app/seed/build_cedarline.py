# -*- coding: utf-8 -*-
"""
Cedarline Outdoor Supply -- deterministic dataset generator
============================================================
MENTOR-ONLY FILE. Reading it reveals every planted data problem.

Builds (relative to the workspace root):
  datasets/sql/cedarline_retail.sql                   PostgreSQL schema + data
  datasets/excel/cedarline_q2_2026_order_lines.xlsx   messy export      (Diagnostic Part B)
  datasets/power_query/cedarline_pos_exports/*.csv    monthly POS files (Diagnostic Part D)
  datasets/power_query/stores.csv
  datasets/power_query/store_targets_summer_2026.csv
  solutions/00_diagnostic/_generated/ground_truth.json

Run:  python solutions/_generators/build_cedarline.py
Same seed -> identical data (verified 2026-09-14: .sql/.csv/.json byte-identical,
.xlsx cell-identical; only the xlsx save timestamp differs). Python 3.8+, openpyxl.
If you change this script, reload the DB: psql -d da_mastery -f datasets/sql/cedarline_retail.sql
"""
import csv
import json
import math
import os
import random
from collections import Counter, defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, os.pardir, os.pardir))
SEED = 20260914
rng = random.Random(SEED)

SIGNUP_START = date(2023, 6, 1)
DATA_START = date(2024, 1, 1)   # order history before this date was not migrated
DATA_END = date(2026, 8, 31)
PRICE_RISE = date(2025, 3, 1)   # 5% list-price increase


def r2(x):
    return round(x + 1e-9, 2)


def month_start(d):
    return date(d.year, d.month, 1)


# --------------------------------------------------------------------------- reference data
SEASON = {1: .65, 2: .70, 3: .90, 4: 1.10, 5: 1.30, 6: 1.40,
          7: 1.35, 8: 1.20, 9: .90, 10: .80, 11: 1.05, 12: 1.25}
SEASON_MAX = 1.40
REGIONS = ['North', 'South', 'East', 'West']

STORES = [  # store_id, name, region, opened_on
    ('N01', 'Northgate', 'North', date(2021, 3, 1)),
    ('S01', 'Riverside', 'South', date(2021, 3, 1)),
    ('E01', 'Harborview', 'East', date(2022, 5, 1)),
    ('W01', 'Summit', 'West', date(2021, 3, 1)),
    ('W02', 'Canyon', 'West', date(2025, 9, 1)),
]
STORE = {s[0]: s for s in STORES}

EMPLOYEES = [  # id, name, title, store, manager, hired, terminated
    (1, 'Dana Whitfield', 'Head of Retail', None, None, date(2020, 11, 2), None),
    (2, 'Marcus Oyelaran', 'Regional Manager - North & East', None, 1, date(2021, 1, 11), None),
    (3, 'Priya Raman', 'Regional Manager - South & West', None, 1, date(2021, 1, 18), None),
    (4, 'Tomasz Nowak', 'E-commerce Manager', None, 1, date(2022, 2, 7), None),
    (5, 'Helen Achebe', 'Store Manager', 'N01', 2, date(2021, 2, 1), None),
    (6, 'Luis Ferreira', 'Store Manager', 'E01', 2, date(2022, 4, 4), None),
    (7, 'Grace Kim', 'Store Manager', 'S01', 3, date(2021, 2, 8), None),
    (8, 'Owen Brandt', 'Store Manager', 'W01', 3, date(2021, 2, 15), None),
    (9, 'Amara Osei', 'Store Manager', 'W02', 3, date(2025, 7, 14), None),
    (10, 'Jake Morrison', 'Sales Associate', 'N01', 5, date(2021, 2, 22), None),
    (11, 'Sofia Petrov', 'Sales Associate', 'N01', 5, date(2021, 6, 7), None),
    (12, 'Daniel Reyes', 'Sales Associate', 'N01', 5, date(2022, 3, 14), None),
    (13, 'Mei Lin', 'Sales Associate', 'N01', 5, date(2022, 9, 5), date(2025, 6, 30)),
    (14, 'Noah Fischer', 'Sales Associate', 'N01', 5, date(2025, 7, 7), None),
    (15, 'Aisha Bello', 'Sales Associate', 'S01', 7, date(2021, 2, 22), None),
    (16, "Ryan O'Connor", 'Sales Associate', 'S01', 7, date(2021, 8, 16), None),
    (17, 'Elena Rossi', 'Sales Associate', 'S01', 7, date(2022, 5, 2), None),
    (18, 'Kwame Mensah', 'Sales Associate', 'S01', 7, date(2023, 4, 10), None),
    (19, 'Chloe Martin', 'Sales Associate', 'E01', 6, date(2022, 4, 18), None),
    (20, 'Arjun Mehta', 'Sales Associate', 'E01', 6, date(2022, 4, 18), None),
    (21, 'Isabel Cruz', 'Sales Associate', 'E01', 6, date(2023, 1, 9), None),
    (22, 'Samuel Park', 'Sales Associate', 'E01', 6, date(2022, 6, 6), date(2024, 10, 31)),
    (23, 'Leah Goldberg', 'Sales Associate', 'E01', 6, date(2024, 11, 4), None),
    (24, 'Ethan Walsh', 'Sales Associate', 'W01', 8, date(2021, 3, 1), None),
    (25, 'Nadia Haddad', 'Sales Associate', 'W01', 8, date(2021, 9, 13), None),
    (26, 'Victor Salas', 'Sales Associate', 'W01', 8, date(2022, 7, 11), None),
    (27, 'Zoe Turner', 'Sales Associate', 'W01', 8, date(2023, 5, 22), None),
    (28, 'Liam Novak', 'Sales Associate', 'W02', 9, date(2025, 8, 18), None),
    (29, 'Fatima Zahra', 'Sales Associate', 'W02', 9, date(2025, 8, 18), None),
    (30, 'Ben Carter', 'Sales Associate', 'W02', 9, date(2025, 8, 25), None),
    (31, 'Hannah Cole', 'E-commerce Analyst', None, 4, date(2023, 3, 6), None),
]
EMP = {e[0]: e for e in EMPLOYEES}

CAT_COST = {'Tents': .58, 'Sleeping Gear': .56, 'Backpacks': .50, 'Footwear': .55,
            'Apparel': .42, 'Cooking': .52, 'Accessories': .40}
CAT_BASE = {'Tents': .08, 'Sleeping Gear': .10, 'Backpacks': .13, 'Footwear': .18,
            'Apparel': .24, 'Cooking': .12, 'Accessories': .15}
CAT_SEASON = {
    'Tents': lambda m: 1.7 if m in (4, 5, 6, 7, 8) else (0.9 if m in (3, 9) else 0.45),
    'Sleeping Gear': lambda m: 1.4 if m in (4, 5, 6, 7, 8) else (1.0 if m in (3, 9, 10) else 0.7),
    'Backpacks': lambda m: 1.3 if m in (7, 8, 9) else 1.0,
    'Footwear': lambda m: 1.2 if m in (3, 4, 5, 9, 10) else 1.0,
    'Apparel': lambda m: 1.6 if m in (10, 11, 12, 1, 2) else (0.75 if m in (5, 6, 7, 8) else 1.0),
    'Cooking': lambda m: 1.4 if m in (5, 6, 7, 8) else 0.8,
    'Accessories': lambda m: 1.5 if m == 12 else (1.2 if m == 11 else 1.0),
}
RET_P = {'Footwear': .12, 'Apparel': .10, 'Backpacks': .05, 'Tents': .06,
         'Sleeping Gear': .05, 'Cooking': .02, 'Accessories': .03}

D0 = date(2021, 3, 1)
NEW = date(2026, 9, 1)
PRODUCTS = [  # sku, name, category, list price, launch, discontinued_on, clearance_from
    ('TNT-100', 'Classic A-Frame Tent', 'Tents', 159, D0, date(2025, 1, 1), None),
    ('TNT-101', 'Ridgeline 1P Tent', 'Tents', 189, D0, None, None),
    ('TNT-102', 'Ridgeline 2P Tent', 'Tents', 249, D0, None, None),
    ('TNT-103', 'Basecamp 4P Dome Tent', 'Tents', 329, D0, None, None),
    ('TNT-104', 'Basecamp 6P Cabin Tent', 'Tents', 449, D0, None, None),
    ('TNT-105', 'Ultralight Trek Tarp', 'Tents', 119, date(2024, 3, 1), None, None),
    ('TNT-106', 'Four-Season Summit Tent', 'Tents', 599, D0, None, None),
    ('TNT-107', 'Family Screen Shelter', 'Tents', 279, D0, None, None),
    ('TNT-108', 'Bivy Sack', 'Tents', 139, date(2025, 4, 1), None, None),
    ('TNT-109', 'Ultralight 2P Trekking Tent', 'Tents', 399, NEW, None, None),
    ('SLP-201', 'Nightfall 30F Sleeping Bag', 'Sleeping Gear', 129, D0, None, None),
    ('SLP-202', 'Nightfall 15F Sleeping Bag', 'Sleeping Gear', 179, D0, None, None),
    ('SLP-203', 'Arctic 0F Down Sleeping Bag', 'Sleeping Gear', 349, D0, None, None),
    ('SLP-204', "Kids' Campfire Sleeping Bag", 'Sleeping Gear', 59, D0, None, None),
    ('SLP-205', 'Self-Inflating Sleeping Pad', 'Sleeping Gear', 89, D0, None, None),
    ('SLP-206', 'Foam Trail Pad', 'Sleeping Gear', 34, D0, None, None),
    ('SLP-207', 'Camp Pillow', 'Sleeping Gear', 24, D0, None, None),
    ('SLP-208', 'Double Air Mattress', 'Sleeping Gear', 119, D0, None, None),
    ('BPK-301', 'Daytripper 22L Pack', 'Backpacks', 69, D0, None, None),
    ('BPK-302', 'Trailhead 35L Pack', 'Backpacks', 129, D0, None, None),
    ('BPK-303', 'Expedition 65L Pack', 'Backpacks', 259, D0, None, None),
    ('BPK-304', 'Hydration Pack 3L', 'Backpacks', 79, D0, None, None),
    ('BPK-305', "Kids' Explorer Pack", 'Backpacks', 39, D0, None, None),
    ('BPK-306', 'Travel Duffel 60L', 'Backpacks', 99, D0, None, None),
    ('BPK-307', 'Hip Pack', 'Backpacks', 29, D0, None, None),
    ('BPK-308', 'Trek 50L Pack', 'Backpacks', 189, D0, date(2026, 6, 1), date(2026, 3, 1)),
    ('FTW-401', 'Trailrunner Low Shoe', 'Footwear', 139, D0, None, None),
    ('FTW-402', 'Summit Mid Hiking Boot', 'Footwear', 189, D0, None, None),
    ('FTW-403', 'Alpine Leather Boot', 'Footwear', 249, D0, None, None),
    ('FTW-404', 'Camp Sandal', 'Footwear', 59, D0, None, None),
    ('FTW-405', 'Winter Insulated Boot', 'Footwear', 219, D0, None, None),
    ('FTW-406', 'Approach Shoe', 'Footwear', 159, D0, None, None),
    ('FTW-407', "Kids' Hiking Boot", 'Footwear', 79, D0, None, None),
    ('FTW-408', 'Water Shoe', 'Footwear', 49, D0, None, date(2026, 5, 1)),
    ('APP-501', 'Storm Shell Rain Jacket', 'Apparel', 169, D0, None, None),
    ('APP-502', 'Down Puffer Jacket', 'Apparel', 229, D0, None, None),
    ('APP-503', 'Fleece Midlayer', 'Apparel', 89, D0, None, None),
    ('APP-504', 'Merino Base Layer Top', 'Apparel', 79, D0, None, None),
    ('APP-505', 'Convertible Hiking Pants', 'Apparel', 69, D0, None, None),
    ('APP-506', 'Sun Hoodie', 'Apparel', 49, D0, None, None),
    ('APP-507', 'Wool Hiking Socks 3-Pack', 'Apparel', 29, D0, None, None),
    ('APP-508', 'Softshell Gloves', 'Apparel', 39, D0, None, None),
    ('APP-509', 'Trail Cap', 'Apparel', 24, D0, None, None),
    ('CKG-601', 'Compact Canister Stove', 'Cooking', 49, D0, None, None),
    ('CKG-602', 'Two-Burner Camp Stove', 'Cooking', 129, D0, None, None),
    ('CKG-603', 'Titanium Cook Pot 1L', 'Cooking', 59, D0, None, None),
    ('CKG-604', 'Camp Cookset 4-Person', 'Cooking', 89, D0, None, None),
    ('CKG-605', 'Insulated Camp Mug', 'Cooking', 25, D0, None, None),
    ('CKG-606', 'Water Filter Bottle', 'Cooking', 45, D0, None, None),
    ('CKG-607', 'Gravity Water Filter', 'Cooking', 99, D0, None, None),
    ('CKG-608', 'Rolling Cooler 50QT', 'Cooking', 159, D0, None, None),
    ('CKG-609', 'Pour-Over Coffee Kit', 'Cooking', 35, NEW, None, None),
    ('ACC-701', 'LED Headlamp 300lm', 'Accessories', 39, D0, None, None),
    ('ACC-702', 'Rechargeable Lantern', 'Accessories', 49, D0, None, None),
    ('ACC-703', 'Trekking Poles (Pair)', 'Accessories', 89, D0, None, None),
    ('ACC-704', 'Multi-Tool', 'Accessories', 59, D0, None, None),
    ('ACC-705', 'First Aid Kit', 'Accessories', 35, D0, None, None),
    ('ACC-706', 'Dry Bag 20L', 'Accessories', 29, D0, None, None),
    ('ACC-707', 'Camp Chair', 'Accessories', 69, D0, None, None),
    ('ACC-708', 'Solar Power Bank', 'Accessories', 79, D0, None, None),
    ('ACC-709', 'Carabiner Set', 'Accessories', 19, NEW, None, None),
]
CONSUMABLE = {'APP-507', 'APP-509', 'CKG-605', 'ACC-706', 'SLP-207', 'ACC-705'}

FIRST = ['Olivia', 'Liam', 'Emma', 'Noah', 'Ava', 'Mateo', 'Sophia', 'Lucas', 'Mia', 'Ethan',
         'Amelia', 'Omar', 'Harper', 'Leo', 'Evelyn', 'Yusuf', 'Aria', 'Daniel', 'Chloe', 'Samuel',
         'Layla', 'Henry', 'Zoe', 'Kenji', 'Nora', 'Diego', 'Hana', 'Isaac', 'Leila', 'Caleb',
         'Priya', 'Aaron', 'Maya', 'Hugo', 'Ines', 'Felix', 'Alice', 'Tariq', 'Clara', 'Ravi',
         'Julia', 'Marco', 'Elif', 'Jonas', 'Grace', 'Adam', 'Sara', 'Andre', 'Lucy', 'Kofi',
         'Emily', 'Arjun', 'Rosa', 'Viktor', 'Naomi', 'Jamal', 'Freya', 'Chen', 'Lina', 'Oscar',
         'Mariam', 'Theo', 'Ruth', 'Pablo', 'Anya', 'Malik', 'Ivy', 'Rahul', 'Esther', 'José']
LAST = ['Smith', 'Johnson', 'Garcia', 'Nguyen', 'Patel', 'Kim', 'Hassan', 'Brown', 'Martinez',
        'Lopez', 'Wilson', 'Anderson', 'Thomas', 'Taylor', 'Moore', 'Jackson', 'Lee', 'Harris',
        'Clark', 'Lewis', 'Walker', 'Young', 'Allen', 'King', 'Wright', 'Scott', 'Green', 'Baker',
        'Adams', 'Nelson', 'Hill', 'Campbell', 'Mitchell', 'Roberts', 'Carter', 'Phillips', 'Evans',
        'Turner', 'Torres', 'Parker', 'Collins', 'Edwards', 'Stewart', 'Morris', 'Murphy', 'Rivera',
        'Cook', 'Rogers', 'Morgan', 'Cooper', 'Reed', 'Bailey', 'Bell', 'Gomez', 'Kelly', 'Ward',
        'Diaz', 'Wood', 'Watson', 'Brooks', 'Bennett', 'Gray', 'Reyes', 'Cruz', 'Hughes', 'Price',
        'Foster', 'Sanders', 'Ross', 'Morales', 'Powell', 'Sullivan', 'Russell', 'Ortiz', 'Perry',
        'Fisher', "O'Brien", "D'Souza", 'Okafor', 'Haddad', 'Kowalski', 'Yilmaz', 'Tanaka', 'Silva',
        'Novak', 'Ivanova', 'Mensah', 'Farouk', 'Rahman', 'Schmidt', 'Dubois', 'Rossi', 'Larsen',
        'Moreno', 'Muñoz', 'Chowdhury', 'Abbasi', 'Lindqvist', 'Mwangi', 'Castillo']
CITIES = {'North': ['Ashford', 'Pinecrest', 'Millbrook', 'Northfield'],
          'South': ['Bayport', 'Marlow', 'Southwick', 'Palmetto Bay'],
          'East': ['Easton', 'Kingsbridge', 'Harbor City', 'Lakewood'],
          'West': ['Redwood Falls', 'Silver Creek', 'Canyon Ridge', 'Westvale']}
DOMAINS = ['mailbox.example', 'trailmail.example', 'inbox.example', 'post.example', 'webmail.example']
ACQ = ['Organic Search', 'Paid Social', 'Referral', 'Email', 'In-Store']
ACQ_W = [.28, .22, .12, .08, .30]
TYPES = {'casual': (1.3, .30), 'regular': (3.2, .12), 'enthusiast': (7.0, .05)}
TYPE_MIX = {'Paid Social': (.75, .20, .05), 'Referral': (.45, .40, .15)}

# --------------------------------------------------------------------------- products
products = []
for pid, (sku, name, cat, price, launch, disc_on, clear_from) in enumerate(PRODUCTS, start=1):
    ratio = CAT_COST[cat] + rng.uniform(-0.04, 0.04)
    products.append(dict(
        product_id=pid, sku=sku, name=name, category=cat, list_price=float(price),
        unit_cost=r2(price * ratio), launch=launch, disc_on=disc_on, clear_from=clear_from,
        is_active=(disc_on is None and clear_from is None),
        pop=rng.uniform(0.6, 1.5) / math.sqrt(price)))
PROD = {p['product_id']: p for p in products}
PROD_BY_SKU = {p['sku']: p for p in products}


def unit_price(p, d):
    if p['clear_from'] and d >= p['clear_from']:
        return r2(p['list_price'] * 0.7)
    if d < PRICE_RISE:
        return float(round(p['list_price'] * 0.95))
    return p['list_price']


def pick_products(d, n):
    avail = [p for p in products if p['launch'] <= d and (p['disc_on'] is None or d < p['disc_on'])]
    cat_pop = defaultdict(float)
    for p in avail:
        cat_pop[p['category']] += p['pop']
    weights = []
    for p in avail:
        w = CAT_BASE[p['category']] * CAT_SEASON[p['category']](d.month) * p['pop'] / cat_pop[p['category']]
        if p['clear_from'] and d >= p['clear_from']:
            w *= 2.5
        weights.append(w)
    chosen = []
    while len(chosen) < n:
        p = rng.choices(avail, weights=weights)[0]
        if p not in chosen:
            chosen.append(p)
    return chosen


# --------------------------------------------------------------------------- people
def strip_accents(s):
    return s.replace('ñ', 'n').replace('é', 'e')


used_emails = set()


def make_email(first, last):
    base = '%s.%s' % (strip_accents(first).lower(),
                      strip_accents(last).lower().replace("'", '').replace(' ', ''))
    dom = rng.choice(DOMAINS)
    email = '%s@%s' % (base, dom)
    while email in used_emails:
        email = '%s%d@%s' % (base, rng.randint(2, 99), dom)
    used_emails.add(email)
    return email


def sample_signup():
    span = (DATA_END - SIGNUP_START).days
    cap = math.exp(0.0006 * span) * SEASON_MAX
    while True:
        k = rng.randrange(span + 1)
        d = SIGNUP_START + timedelta(days=k)
        if rng.random() < math.exp(0.0006 * k) * SEASON[d.month] / cap:
            return d


N_PERSONS = 4500
persons = []
for k in range(N_PERSONS):
    first, last = rng.choice(FIRST), rng.choice(LAST)
    region = None if rng.random() < 0.02 else rng.choices(REGIONS, weights=[.27, .24, .21, .28])[0]
    city = rng.choice(CITIES[region] if region else CITIES[rng.choice(REGIONS)])
    acq = None if rng.random() < 0.05 else rng.choices(ACQ, weights=ACQ_W)[0]
    mix = TYPE_MIX.get(acq, (.60, .30, .10))
    ptype = rng.choices(['casual', 'regular', 'enthusiast'], weights=mix)[0]
    p_online = 0.15 if acq == 'In-Store' else (0.45 if acq is None else 0.70)
    persons.append(dict(pid=k, first=first, last=last, name=first + ' ' + last,
                        email=make_email(first, last), city=city, region=region,
                        signup=sample_signup(), acq=acq, ptype=ptype, p_online=p_online,
                        never=rng.random() < (0.18 if acq == 'Paid Social' else 0.12)))


def pick_store(region, d):
    open_ids = [s[0] for s in STORES if s[3] <= d]
    if region is None or rng.random() < 0.10:
        return rng.choice(open_ids)
    if region == 'West':
        if d >= STORE['W02'][3]:
            months = (d.year - 2025) * 12 + d.month - 9
            share = min(0.45, 0.20 + 0.04 * months)
            return 'W02' if rng.random() < share else 'W01'
        return 'W01'
    return {'North': 'N01', 'South': 'S01', 'East': 'E01'}[region]


def pick_employee(store_id, d):
    staff = [e for e in EMPLOYEES if e[3] == store_id and e[5] <= d and (e[6] is None or e[6] >= d)]
    assoc = [e for e in staff if e[2] == 'Sales Associate']
    pool = assoc if assoc and rng.random() < 0.85 else staff
    return rng.choice(pool)[0]


# --------------------------------------------------------------------------- order simulation
orders = []  # dicts; customer record assigned later
for p in persons:
    if p['never']:
        continue
    rate, churn = TYPES[p['ptype']]
    rate_max = rate / 0.73
    lag = 0 if (p['acq'] == 'In-Store' and rng.random() < 0.75) else int(rng.expovariate(1 / 15.0))
    t = p['signup'] + timedelta(days=lag)
    first_order = True
    while t <= DATA_END:
        if first_order or rng.random() < SEASON[t.month] / SEASON_MAX:
            first_order = False
            online = rng.random() < p['p_online']
            if t >= DATA_START:
                orders.append(dict(pid=p['pid'], date=t, channel='Online' if online else 'In-Store',
                                   store_id=None if online else pick_store(p['region'], t)))
            if rng.random() < churn:
                break
        t = t + timedelta(days=max(1, int(rng.expovariate(rate_max / 365.0))))

# walk-in (anonymous) in-store orders: ~9% on top of identified in-store orders
instore_counts = Counter((o['store_id'], month_start(o['date'])) for o in orders if o['channel'] == 'In-Store')
for (sid, ms), n in sorted(instore_counts.items()):
    n_walk = int(round(0.09 * n * rng.uniform(0.8, 1.2)))
    last_day = (month_start(ms + timedelta(days=32)) - timedelta(days=1))
    last_day = min(last_day, DATA_END)
    first_day = max(ms, STORE[sid][3])
    for _ in range(n_walk):
        d = first_day + timedelta(days=rng.randrange((last_day - first_day).days + 1))
        orders.append(dict(pid=None, date=d, channel='In-Store', store_id=sid))

# --------------------------------------------------------------------------- customer records (+ duplicates)
by_person = defaultdict(list)
for o in orders:
    if o['pid'] is not None:
        by_person[o['pid']].append(o)
for lst in by_person.values():
    lst.sort(key=lambda o: o['date'])

records = []  # customer records; 'rec' index assigned in order of creation
for p in persons:
    p['recs'] = [len(records)]
    records.append(dict(pid=p['pid'], name=p['name'], email=p['email'], city=p['city'],
                        region=p['region'], signup=p['signup'], acq=p['acq']))


def email_variant(e):
    local, dom = e.split('@')
    choice = rng.randrange(4)
    if choice == 0:
        return '.'.join(w.capitalize() for w in local.split('.')) + '@' + dom.capitalize()
    if choice == 1:
        return e.upper()
    if choice == 2:
        return ' ' + e
    return e + ' '


def name_variant(n):
    choice = rng.randrange(3)
    return n if choice == 0 else (n.upper() if choice == 1 else n.lower())


def add_dup_record(p, signup):
    base = records[p['recs'][0]]
    records.append(dict(pid=p['pid'], name=name_variant(base['name']), email=email_variant(base['email']),
                        city=base['city'], region=base['region'], signup=signup,
                        acq=base['acq'] if rng.random() < 0.5 else None))
    p['recs'].append(len(records) - 1)
    return len(records) - 1


multi = [p for p in persons if len(by_person[p['pid']]) >= 3]
dup_people = rng.sample(multi, 60)
order_rec = {}
for i, p in enumerate(dup_people):
    olist = by_person[p['pid']]
    j = rng.randrange(1, len(olist))
    rec = add_dup_record(p, olist[j]['date'])
    for o in olist[j:]:
        order_rec[id(o)] = rec if (o is olist[j] or rng.random() < 0.6) else p['recs'][0]
    if i < 4 and j + 1 < len(olist):  # a third record for four people
        j2 = rng.randrange(j + 1, len(olist))
        rec3 = add_dup_record(p, olist[j2]['date'])
        for o in olist[j2:]:
            if o is olist[j2] or rng.random() < 0.5:
                order_rec[id(o)] = rec3
light = [p for p in persons if len(by_person[p['pid']]) <= 1 and p not in dup_people]
for p in rng.sample(light, 8):
    add_dup_record(p, p['signup'] + timedelta(days=rng.randint(20, 200)))

# assign customer_id by signup order
rec_order = sorted(range(len(records)), key=lambda r: (records[r]['signup'], rng.random()))
for n, r in enumerate(rec_order):
    records[r]['customer_id'] = 10001 + n
for r in records:
    if r['signup'] > DATA_END:
        r['signup'] = DATA_END

for o in orders:
    if o['pid'] is None:
        o['customer_id'] = None
    else:
        rec = order_rec.get(id(o), persons[o['pid']]['recs'][0])
        o['customer_id'] = records[rec]['customer_id']
CUST = {r['customer_id']: r for r in records}

# header-only orders (payment failures etc.) -- no order_items
for k in range(12):
    d = DATA_START + timedelta(days=rng.randrange((DATA_END - DATA_START).days + 1))
    r = rng.choice(records)
    orders.append(dict(pid=r['pid'], customer_id=r['customer_id'], date=d, channel='Online',
                       store_id=None, header_only=True, status='cancelled' if k < 8 else 'completed'))

orders.sort(key=lambda o: (o['date'], rng.random()))
for n, o in enumerate(orders):
    o['order_id'] = 100001 + n

# --------------------------------------------------------------------------- items, status, shipping
items = []
next_item = 500001
for o in orders:
    d = o['date']
    o['employee_id'] = pick_employee(o['store_id'], d) if o['channel'] == 'In-Store' else None
    if o.get('header_only'):
        o['shipping_fee'] = 7.99 if o['status'] == 'completed' else 0.0
        o['items'] = []
        continue
    if o['channel'] == 'Online':
        if d >= date(2026, 8, 27) and rng.random() < 0.45:
            o['status'] = 'pending'
        else:
            o['status'] = 'cancelled' if rng.random() < 0.05 else 'completed'
    else:
        o['status'] = 'cancelled' if rng.random() < 0.02 else 'completed'
    n_lines = rng.choices([1, 2, 3, 4], weights=[.55, .27, .12, .06])[0]
    black_friday = (d.month == 11 and d.day >= 24) or (d.month == 12 and d.day <= 1)
    order_disc = None
    if black_friday and rng.random() < 0.55:
        order_disc = rng.choice([0.20, 0.25])
    elif rng.random() < 0.12:
        order_disc = rng.choice([0.10, 0.15, 0.15, 0.20])
    o['items'] = []
    for p in pick_products(d, n_lines):
        if p['sku'] in CONSUMABLE:
            qty = rng.choices([1, 2, 3, 4], weights=[.5, .3, .12, .08])[0]
        else:
            qty = rng.choices([1, 2, 3], weights=[.86, .11, .03])[0]
        price = unit_price(p, d)
        if p['clear_from'] and d >= p['clear_from']:
            disc = 0.0
        elif order_disc is not None:
            disc = order_disc
        else:
            u = rng.random()
            disc = 0.10 if u < 0.08 else (0.15 if u < 0.11 else 0.0)
        it = dict(order_item_id=next_item, order_id=o['order_id'], product_id=p['product_id'],
                  quantity=qty, unit_price=price, discount_pct=disc)
        next_item += 1
        items.append(it)
        o['items'].append(it)
    subtotal = sum(i['quantity'] * i['unit_price'] * (1 - i['discount_pct']) for i in o['items'])
    o['shipping_fee'] = (0.0 if subtotal >= 100 else 7.99) if o['channel'] == 'Online' else 0.0
ORDER = {o['order_id']: o for o in orders}

# --------------------------------------------------------------------------- returns
REASONS_SIZE = (['Wrong size', 'Changed mind', 'Defective', 'Not as described', 'Damaged in transit'],
                [.55, .20, .10, .10, .05])
REASONS_OTHER = (['Defective', 'Changed mind', 'Not as described', 'Damaged in transit'],
                 [.35, .30, .20, .15])
returns = []
for it in items:
    o = ORDER[it['order_id']]
    if o['status'] != 'completed':
        continue
    p = PROD[it['product_id']]
    prob = RET_P[p['category']] * (1.5 if o['channel'] == 'Online' else 0.7)
    if rng.random() >= prob:
        continue
    horizon = 30 if o['channel'] == 'In-Store' else 45
    rdate = o['date'] + timedelta(days=rng.randint(2, horizon))
    if rdate > DATA_END:
        continue
    pool, w = REASONS_SIZE if p['category'] in ('Footwear', 'Apparel') else REASONS_OTHER
    reason = rng.choices(pool, weights=w)[0]
    if reason == 'Damaged in transit' and o['channel'] == 'In-Store':
        reason = 'Changed mind'
    q = it['quantity']
    q1 = q if (q == 1 or rng.random() < 0.7) else rng.randint(1, q - 1)
    net_unit = it['unit_price'] * (1 - it['discount_pct'])
    returns.append(dict(order_item_id=it['order_item_id'], return_date=rdate, quantity_returned=q1,
                        reason=reason, refund_amount=r2(q1 * net_unit)))
    if q1 < q and rng.random() < 0.35:
        rdate2 = rdate + timedelta(days=rng.randint(3, 20))
        if rdate2 <= DATA_END:
            returns.append(dict(order_item_id=it['order_item_id'], return_date=rdate2,
                                quantity_returned=q - q1, reason=reason,
                                refund_amount=r2((q - q1) * net_unit)))
returns.sort(key=lambda r: (r['return_date'], r['order_item_id'], rng.random()))
for n, r in enumerate(returns):
    r['return_id'] = 900001 + n
RETURNS_BY_ITEM = defaultdict(list)
for r in returns:
    RETURNS_BY_ITEM[r['order_item_id']].append(r)


def order_region(o):
    if o['channel'] == 'In-Store':
        return STORE[o['store_id']][2]
    return CUST[o['customer_id']]['region'] if o['customer_id'] else None


def line_net(it):
    return it['quantity'] * it['unit_price'] * (1 - it['discount_pct'])


# =========================================================================== SQL FILE
def sql_val(v):
    if v is None:
        return 'NULL'
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return '%.2f' % v
    if isinstance(v, date):
        return "'%s'" % v.isoformat()
    return "'" + str(v).replace("'", "''") + "'"


def inserts(table, cols, rows, batch=1000):
    out = []
    for i in range(0, len(rows), batch):
        chunk = rows[i:i + batch]
        out.append('INSERT INTO %s (%s) VALUES\n' % (table, ', '.join(cols)) +
                   ',\n'.join('(' + ', '.join(sql_val(v) for v in r) + ')' for r in chunk) + ';\n')
    return '\n'.join(out)


DDL = """-- =====================================================================
-- Cedarline Outdoor Supply -- practice database
-- Dialect : PostgreSQL (tested on PostgreSQL 18)
-- Load    : psql -d da_mastery -v ON_ERROR_STOP=1 -f cedarline_retail.sql
-- Re-running this file drops and rebuilds the whole "cedarline" schema,
-- so it doubles as a reset button if you break something.
-- Data window: orders 2024-01-01 .. 2026-08-31 (older history not migrated)
-- =====================================================================
SET client_encoding = 'UTF8';
DROP SCHEMA IF EXISTS cedarline CASCADE;
CREATE SCHEMA cedarline;
SET search_path TO cedarline;

CREATE TABLE stores (
    store_id    VARCHAR(3)  PRIMARY KEY,
    store_name  VARCHAR(40) NOT NULL,
    region      VARCHAR(10) NOT NULL,
    opened_on   DATE        NOT NULL
);

CREATE TABLE employees (
    employee_id      INTEGER     PRIMARY KEY,
    full_name        VARCHAR(60) NOT NULL,
    job_title        VARCHAR(60) NOT NULL,
    store_id         VARCHAR(3)  REFERENCES stores (store_id),
    manager_id       INTEGER     REFERENCES employees (employee_id),
    hire_date        DATE        NOT NULL,
    termination_date DATE
);

CREATE TABLE customers (
    customer_id         INTEGER      PRIMARY KEY,
    full_name           VARCHAR(80)  NOT NULL,
    email               VARCHAR(120) NOT NULL,
    city                VARCHAR(40),
    region              VARCHAR(10),
    signup_date         DATE         NOT NULL,
    acquisition_channel VARCHAR(20)
);

CREATE TABLE products (
    product_id      INTEGER       PRIMARY KEY,
    sku             VARCHAR(10)   NOT NULL UNIQUE,
    product_name    VARCHAR(60)   NOT NULL,
    category        VARCHAR(20)   NOT NULL,
    unit_cost       NUMERIC(10,2) NOT NULL,
    list_price      NUMERIC(10,2) NOT NULL,
    launch_date     DATE          NOT NULL,
    discontinued_on DATE,
    is_active       BOOLEAN       NOT NULL
);

CREATE TABLE orders (
    order_id     INTEGER      PRIMARY KEY,
    customer_id  INTEGER      REFERENCES customers (customer_id),
    order_date   DATE         NOT NULL,
    channel      VARCHAR(10)  NOT NULL CHECK (channel IN ('Online', 'In-Store')),
    store_id     VARCHAR(3)   REFERENCES stores (store_id),
    employee_id  INTEGER      REFERENCES employees (employee_id),
    status       VARCHAR(10)  NOT NULL CHECK (status IN ('completed', 'cancelled', 'pending')),
    shipping_fee NUMERIC(8,2) NOT NULL DEFAULT 0
);

CREATE TABLE order_items (
    order_item_id INTEGER       PRIMARY KEY,
    order_id      INTEGER       NOT NULL REFERENCES orders (order_id),
    product_id    INTEGER       NOT NULL REFERENCES products (product_id),
    quantity      INTEGER       NOT NULL CHECK (quantity > 0),
    unit_price    NUMERIC(10,2) NOT NULL CHECK (unit_price >= 0),
    discount_pct  NUMERIC(4,2)  NOT NULL DEFAULT 0 CHECK (discount_pct >= 0 AND discount_pct < 1)
);

CREATE TABLE returns (
    return_id         INTEGER       PRIMARY KEY,
    order_item_id     INTEGER       NOT NULL REFERENCES order_items (order_item_id),
    return_date       DATE          NOT NULL,
    quantity_returned INTEGER       NOT NULL CHECK (quantity_returned > 0),
    reason            VARCHAR(30)   NOT NULL,
    refund_amount     NUMERIC(10,2) NOT NULL
);

"""

COMMENTS = """
CREATE INDEX idx_orders_date      ON orders (order_date);
CREATE INDEX idx_orders_customer  ON orders (customer_id);
CREATE INDEX idx_items_order      ON order_items (order_id);
CREATE INDEX idx_items_product    ON order_items (product_id);
CREATE INDEX idx_returns_item     ON returns (order_item_id);

COMMENT ON TABLE  orders IS 'One row per order (header). Order history starts 2024-01-01.';
COMMENT ON COLUMN orders.customer_id  IS 'NULL = anonymous walk-in purchase in a store.';
COMMENT ON COLUMN orders.store_id     IS 'Store that made the sale; NULL for Online orders.';
COMMENT ON COLUMN orders.employee_id  IS 'Associate who rang the sale; NULL for Online orders.';
COMMENT ON COLUMN orders.shipping_fee IS 'Charged once per order (Online only).';
COMMENT ON TABLE  order_items IS 'One row per product line on an order.';
COMMENT ON COLUMN order_items.unit_price   IS 'Price actually charged per unit at the time of sale (before line discount).';
COMMENT ON COLUMN order_items.discount_pct IS 'Line discount as a fraction: 0.15 = 15% off.';
COMMENT ON TABLE  returns IS 'One row per return event. A line can be returned in more than one event.';
COMMENT ON COLUMN products.list_price IS 'CURRENT list price. Historical selling prices live in order_items.unit_price.';
COMMENT ON COLUMN products.is_active  IS 'FALSE = being phased out / no longer restocked.';
COMMENT ON COLUMN customers.region    IS 'Customer home region. Can be NULL.';
"""

sql_parts = [DDL]
sql_parts.append(inserts('stores', ['store_id', 'store_name', 'region', 'opened_on'], [list(s) for s in STORES]))
sql_parts.append(inserts('employees', ['employee_id', 'full_name', 'job_title', 'store_id', 'manager_id',
                                       'hire_date', 'termination_date'], [list(e) for e in EMPLOYEES]))
sql_parts.append(inserts('customers', ['customer_id', 'full_name', 'email', 'city', 'region', 'signup_date',
                                       'acquisition_channel'],
                         [[r['customer_id'], r['name'], r['email'], r['city'], r['region'], r['signup'], r['acq']]
                          for r in sorted(records, key=lambda r: r['customer_id'])]))
sql_parts.append(inserts('products', ['product_id', 'sku', 'product_name', 'category', 'unit_cost', 'list_price',
                                      'launch_date', 'discontinued_on', 'is_active'],
                         [[p['product_id'], p['sku'], p['name'], p['category'], p['unit_cost'], p['list_price'],
                           p['launch'], p['disc_on'], p['is_active']] for p in products]))
sql_parts.append(inserts('orders', ['order_id', 'customer_id', 'order_date', 'channel', 'store_id', 'employee_id',
                                    'status', 'shipping_fee'],
                         [[o['order_id'], o['customer_id'], o['date'], o['channel'], o['store_id'],
                           o['employee_id'], o['status'], o['shipping_fee']] for o in orders]))
sql_parts.append(inserts('order_items', ['order_item_id', 'order_id', 'product_id', 'quantity', 'unit_price',
                                         'discount_pct'],
                         [[i['order_item_id'], i['order_id'], i['product_id'], i['quantity'], i['unit_price'],
                           i['discount_pct']] for i in items]))
sql_parts.append(inserts('returns', ['return_id', 'order_item_id', 'return_date', 'quantity_returned', 'reason',
                                     'refund_amount'],
                         [[r['return_id'], r['order_item_id'], r['return_date'], r['quantity_returned'],
                           r['reason'], r['refund_amount']] for r in returns]))
sql_parts.append(COMMENTS)
sql_path = os.path.join(ROOT, 'datasets', 'sql', 'cedarline_retail.sql')
with open(sql_path, 'w', encoding='utf-8', newline='\n') as f:
    f.write('\n'.join(sql_parts))

# =========================================================================== EXCEL EXPORT (Part B)
Q2_START, Q2_END = date(2026, 4, 1), date(2026, 6, 30)
COLS = ['LineID', 'OrderID', 'OrderDate', 'CustomerID', 'CustomerName', 'Region', 'Channel', 'StoreCode',
        'SalesRep', 'SKU', 'Quantity', 'UnitPrice', 'Discount', 'Status']
truth = []
for o in orders:
    if not (Q2_START <= o['date'] <= Q2_END) or o['status'] == 'pending':
        continue
    for it in o['items']:
        p = PROD[it['product_id']]
        cust = CUST.get(o['customer_id'])
        truth.append(dict(
            LineID=it['order_item_id'], OrderID=o['order_id'], OrderDate=o['date'],
            CustomerID=o['customer_id'], CustomerName=cust['name'] if cust else 'Walk-in',
            Region=order_region(o) or '', Channel=o['channel'], StoreCode=o['store_id'] or '',
            SalesRep=EMP[o['employee_id']][1] if o['employee_id'] else '', SKU=p['sku'],
            Quantity=it['quantity'], UnitPrice=it['unit_price'], Discount=it['discount_pct'],
            Status='Completed' if o['status'] == 'completed' else 'Cancelled',
            category=p['category'], unit_cost=p['unit_cost'], list_price=p['list_price']))
truth.sort(key=lambda l: (l['OrderDate'], l['LineID']))
for l in truth:
    l['net'] = l['Quantity'] * l['UnitPrice'] * (1 - l['Discount'])

# top customers (truth) -- used to keep the "lost CustomerID" plant away from the leaderboard
cust_net = defaultdict(float)
for l in truth:
    if l['Status'] == 'Completed' and l['CustomerID']:
        cust_net[l['CustomerID']] += l['net']
top10 = set(c for c, _ in sorted(cust_net.items(), key=lambda kv: -kv[1])[:10])

raw = [dict((c, l[c]) for c in COLS) for l in truth]
used = set()
issues = defaultdict(list)  # issue -> list of truth indexes


def take(name, k, pred):
    cands = [i for i in range(len(truth)) if i not in used and pred(truth[i])]
    chosen = sorted(rng.sample(cands, k))
    used.update(chosen)
    issues[name] = chosen
    return chosen


REGION_VARIANTS = {'North': ['north', 'NORTH', 'North ', ' North', 'N.'],
                   'South': ['south', 'SOUTH', 'South ', ' South', 'Sth'],
                   'East': ['east', 'EAST', 'East ', ' East', 'E.'],
                   'West': ['west', 'WEST', 'West ', ' West', 'W.']}
for i in take('region_variant', 120, lambda l: l['Region'] != ''):
    raw[i]['Region'] = rng.choice(REGION_VARIANTS[truth[i]['Region']])
for i in take('customer_name_spacing_case', 90, lambda l: l['CustomerID'] is not None):
    n = truth[i]['CustomerName']
    raw[i]['CustomerName'] = rng.choice(['  ' + n, n + '  ', n.upper(), n.lower(), n.replace(' ', '   ', 1)])
for i in take('date_as_text_iso', 45, lambda l: True):
    raw[i]['OrderDate'] = truth[i]['OrderDate'].isoformat()
for i in take('date_as_text_dotted', 10, lambda l: True):
    d = truth[i]['OrderDate']
    raw[i]['OrderDate'] = '%02d.%02d.%04d' % (d.day, d.month, d.year)
for i in take('quantity_as_text', 20, lambda l: True):
    raw[i]['Quantity'] = str(truth[i]['Quantity'])
for i in take('customer_id_lost', 25, lambda l: l['CustomerID'] is not None and l['CustomerID'] not in top10):
    raw[i]['CustomerID'] = None
for i in take('sku_trailing_space', 12, lambda l: l['SKU'] != 'BPK-308'):
    raw[i]['SKU'] = truth[i]['SKU'] + ' '
for i in take('discount_whole_percent', 35, lambda l: l['Discount'] > 0):
    raw[i]['Discount'] = int(round(truth[i]['Discount'] * 100))
for i in take('discount_blank', 40, lambda l: l['Discount'] == 0):
    raw[i]['Discount'] = None
for i in take('unit_price_x100_typo', 5, lambda l: l['Status'] == 'Completed' and l['SKU'] != 'BPK-308'):
    raw[i]['UnitPrice'] = r2(truth[i]['UnitPrice'] * 100)
for i in take('status_spelling_canceled', 30, lambda l: l['Status'] == 'Cancelled'):
    raw[i]['Status'] = 'Canceled'
bpk308_idx = [i for i, l in enumerate(truth) if l['SKU'] == 'BPK-308']

# physical layout: (sort_key, kind, payload)
layout = [(float(i), 'line', i) for i in range(len(truth))]
dup_src = take('exact_duplicate_row', 14, lambda l: l['SKU'] != 'BPK-308')
for n, i in enumerate(dup_src):
    key = i + 0.5 if n < 8 else min(len(truth) - 1, i + rng.randint(5, 300)) + 0.25
    layout.append((key, 'dup', i))
neg_src = take('negative_qty_return_row', 8, lambda l: l['Status'] == 'Completed' and l['SKU'] != 'BPK-308' and any(
    r['return_date'] <= Q2_END for r in RETURNS_BY_ITEM[l['LineID']]))
neg_rows = []
for i in neg_src:
    ret = [r for r in RETURNS_BY_ITEM[truth[i]['LineID']] if r['return_date'] <= Q2_END][0]
    row = dict(raw[i])
    row.update(LineID='%d-R' % truth[i]['LineID'], OrderDate=ret['return_date'],
               Quantity=-ret['quantity_returned'], CustomerID=truth[i]['CustomerID'],
               CustomerName=truth[i]['CustomerName'], Region=truth[i]['Region'], SKU=truth[i]['SKU'],
               UnitPrice=truth[i]['UnitPrice'], Discount=truth[i]['Discount'], Status='Completed')
    pos = max(k for k, l in enumerate(truth) if l['OrderDate'] <= ret['return_date'])
    layout.append((pos + 0.75, 'neg', len(neg_rows)))
    neg_rows.append(dict(row=row, truth_index=i, net=-ret['quantity_returned'] * truth[i]['UnitPrice'] *
                         (1 - truth[i]['Discount'])))
for _ in range(2):
    layout.append((rng.randint(300, len(truth) - 300) + 0.9, 'blank', None))
layout.sort(key=lambda t: t[0])

HEADER_ROW = 4
physical = []
for key, kind, payload in layout:
    if kind == 'line' or kind == 'dup':
        physical.append((kind, raw[payload], payload))
    elif kind == 'neg':
        physical.append((kind, neg_rows[payload]['row'], neg_rows[payload]['truth_index']))
    else:
        physical.append(('blank', None, None))
qty_numeric_sum = sum(r['Quantity'] for k, r, _ in physical if k != 'blank' and isinstance(r['Quantity'], int))
physical.append(('total', None, None))

wb = Workbook()
ws = wb.active
ws.title = 'Orders_Export'
base_font = Font(name='Arial', size=10)
ws['A1'] = 'Cedarline Outdoor Supply - Order Lines Export'
ws['A1'].font = Font(name='Arial', size=12, bold=True)
ws['A2'] = 'Period: 2026-04-01 to 2026-06-30 | Source: POS + Web | Exported 2026-07-02 08:14'
ws['A2'].font = Font(name='Arial', size=9, italic=True)
for c, name in enumerate(COLS, start=1):
    cell = ws.cell(row=HEADER_ROW, column=c, value=name)
    cell.font = Font(name='Arial', size=10, bold=True)
row_of = defaultdict(list)  # truth index -> excel rows (incl. dup copies)
excel_rows = {'duplicate_copies': [], 'negative_rows': [], 'blank_rows': [], 'total_row': None}
for n, (kind, rowdict, ti) in enumerate(physical):
    er = HEADER_ROW + 1 + n
    if kind == 'blank':
        excel_rows['blank_rows'].append(er)
        continue
    if kind == 'total':
        ws.cell(row=er, column=1, value='TOTAL').font = Font(name='Arial', size=10, bold=True)
        ws.cell(row=er, column=COLS.index('Quantity') + 1, value=qty_numeric_sum).font = base_font
        excel_rows['total_row'] = er
        continue
    if kind == 'dup':
        excel_rows['duplicate_copies'].append(er)
    elif kind == 'neg':
        excel_rows['negative_rows'].append(er)
    else:
        row_of[ti].append(er)
    for c, col in enumerate(COLS, start=1):
        v = rowdict[col]
        cell = ws.cell(row=er, column=c, value=(None if v == '' else v))
        cell.font = base_font
        if col == 'OrderDate' and isinstance(v, date):
            cell.number_format = 'yyyy-mm-dd'
        elif col in ('UnitPrice',) and isinstance(v, float):
            cell.number_format = '0.00'
        elif col == 'Discount' and isinstance(v, float):
            cell.number_format = '0.00'
widths = [10, 10, 12, 11, 24, 9, 9, 10, 16, 10, 9, 10, 9, 11]
for c, w in enumerate(widths, start=1):
    ws.column_dimensions[chr(64 + c)].width = w
last_row = HEADER_ROW + len(physical)

# Products sheet (master data team deleted discontinued SKUs, incl. BPK-308 and TNT-100)
wp = wb.create_sheet('Products')
pcols = ['SKU', 'ProductName', 'Category', 'UnitCost', 'ListPrice', 'Active']
for c, name in enumerate(pcols, start=1):
    wp.cell(row=1, column=c, value=name).font = Font(name='Arial', size=10, bold=True)
prow = 2
for p in products:
    if p['disc_on'] is not None:
        continue
    vals = [p['sku'], p['name'], p['category'], p['unit_cost'], p['list_price'], 'Y' if p['is_active'] else 'N']
    for c, v in enumerate(vals, start=1):
        cell = wp.cell(row=prow, column=c, value=v)
        cell.font = base_font
        if c in (4, 5):
            cell.number_format = '0.00'
    prow += 1
for c, w in enumerate([10, 30, 15, 10, 10, 8], start=1):
    wp.column_dimensions[chr(64 + c)].width = w

# Targets sheet (wide)
completed = [l for l in truth if l['Status'] == 'Completed']
reg_month = defaultdict(float)
for l in completed:
    reg_month[(l['Region'] or '(blank)', month_start(l['OrderDate']))] += l['net']
MONTHS_Q2 = [date(2026, 4, 1), date(2026, 5, 1), date(2026, 6, 1)]
TARGET_FACTOR = {'North': 0.97, 'South': 1.02, 'East': 0.94, 'West': 1.14}
targets = {}
for reg in REGIONS:
    for m in MONTHS_Q2:
        targets[(reg, m)] = round(reg_month[(reg, m)] * TARGET_FACTOR[reg] * rng.uniform(0.98, 1.02) / 500.0) * 500
wt = wb.create_sheet('Targets')
wt['A1'] = ('Q2 2026 regional targets - net merchandise revenue '
            '(completed orders, after line discounts, before returns, excl. shipping)')
wt['A1'].font = Font(name='Arial', size=10, bold=True)
wt.cell(row=3, column=1, value='Region').font = Font(name='Arial', size=10, bold=True)
for c, m in enumerate(MONTHS_Q2, start=2):
    cell = wt.cell(row=3, column=c, value=m)
    cell.number_format = 'mmm-yyyy'
    cell.font = Font(name='Arial', size=10, bold=True)
for r, reg in enumerate(REGIONS, start=4):
    wt.cell(row=r, column=1, value=reg).font = base_font
    for c, m in enumerate(MONTHS_Q2, start=2):
        cell = wt.cell(row=r, column=c, value=targets[(reg, m)])
        cell.number_format = '#,##0'
        cell.font = base_font
wt.column_dimensions['A'].width = 12
for col in 'BCD':
    wt.column_dimensions[col].width = 12

# Data dictionary sheet
wd = wb.create_sheet('Data_Dictionary')
dd = [
    ('Sheet', 'Column', 'Meaning (as documented by the Sales Ops team)'),
    ('Orders_Export', 'LineID', 'Unique ID of an order line (one product on one order).'),
    ('Orders_Export', 'OrderID', 'Order the line belongs to. An order can have several lines.'),
    ('Orders_Export', 'OrderDate', 'Date the order was placed.'),
    ('Orders_Export', 'CustomerID', 'Loyalty customer ID. Walk-in (anonymous) store sales have no ID.'),
    ('Orders_Export', 'CustomerName', 'Customer name, or "Walk-in" for anonymous store sales.'),
    ('Orders_Export', 'Region', 'Region the sale is credited to: store region for In-Store sales, '
                                'customer home region for Online sales.'),
    ('Orders_Export', 'Channel', 'In-Store or Online.'),
    ('Orders_Export', 'StoreCode', 'Store that made the sale (blank for Online).'),
    ('Orders_Export', 'SalesRep', 'Associate who rang the sale (blank for Online).'),
    ('Orders_Export', 'SKU', 'Product code. Join key to Products[SKU].'),
    ('Orders_Export', 'Quantity', 'Units sold on the line.'),
    ('Orders_Export', 'UnitPrice', 'Price charged per unit, before the line discount.'),
    ('Orders_Export', 'Discount', 'Line discount as a fraction of price: 0.15 = 15% off.'),
    ('Orders_Export', 'Status', 'Completed or Cancelled. Cancelled orders are not revenue.'),
    ('Products', 'UnitCost', 'Current landed cost per unit.'),
    ('Products', 'ListPrice', 'Current list price per unit.'),
    ('Products', 'Active', 'Y = normal range; N = being phased out (clearance pricing possible).'),
    ('Targets', '(months)', 'Monthly net merchandise revenue target per region.'),
    ('Definitions', 'Net revenue', 'Quantity x UnitPrice x (1 - Discount), completed orders only, '
                                   'before returns, excluding shipping.'),
]
for r, row in enumerate(dd, start=1):
    for c, v in enumerate(row, start=1):
        cell = wd.cell(row=r, column=c, value=v)
        cell.font = Font(name='Arial', size=10, bold=(r == 1))
        cell.alignment = Alignment(wrap_text=True, vertical='top')
wd.column_dimensions['A'].width = 16
wd.column_dimensions['B'].width = 14
wd.column_dimensions['C'].width = 90
xlsx_path = os.path.join(ROOT, 'datasets', 'excel', 'cedarline_q2_2026_order_lines.xlsx')
wb.save(xlsx_path)

# ----- Excel ground truth
pop_b2 = [l['Region'] or '(blank)' for l in truth] + [truth[n['truth_index']]['Region'] or '(blank)' for n in neg_rows]
products_sheet_skus = set(p['sku'] for p in products if p['disc_on'] is None)
completed_known = [l for l in completed if l['SKU'] in products_sheet_skus]
cat = defaultdict(lambda: [0.0, 0.0, 0])
for l in completed_known:
    cat[l['category']][0] += l['net']
    cat[l['category']][1] += l['Quantity'] * l['unit_cost']
    cat[l['category']][2] += 1
margin = {k: dict(net=r2(v[0]), cost=r2(v[1]), margin_pct=round((v[0] - v[1]) / v[0] * 100, 2), lines=v[2])
          for k, v in sorted(cat.items())}
typo_impact = sum(truth[i]['Quantity'] * (raw[i]['UnitPrice'] - truth[i]['UnitPrice']) * (1 - truth[i]['Discount'])
                  for i in issues['unit_price_x100_typo'])
cust_net_typo = defaultdict(float, cust_net)
for i in issues['unit_price_x100_typo']:
    if truth[i]['CustomerID']:
        cust_net_typo[truth[i]['CustomerID']] += truth[i]['Quantity'] * (raw[i]['UnitPrice'] - truth[i]['UnitPrice']) * (1 - truth[i]['Discount'])


def top5(d):
    return [dict(customer_id=c, name=CUST[c]['name'], net=r2(v))
            for c, v in sorted(d.items(), key=lambda kv: -kv[1])[:5]]


naive = 0.0
for kind, r, _ in physical:
    if kind in ('blank', 'total'):
        continue
    q = float(r['Quantity'])
    dsc = r['Discount'] if r['Discount'] is not None else 0
    naive += q * r['UnitPrice'] * (1 - dsc)

excel_gt = dict(
    file='datasets/excel/cedarline_q2_2026_order_lines.xlsx',
    header_row=HEADER_ROW, first_data_row=HEADER_ROW + 1, last_row=last_row,
    physical_rows_below_header=len(physical),
    truth_lines=len(truth), truth_orders=len(set(l['OrderID'] for l in truth)),
    unique_lineids_in_file=len(truth) + len(neg_rows),
    duplicate_copy_rows=excel_rows['duplicate_copies'], blank_rows=excel_rows['blank_rows'],
    total_row=excel_rows['total_row'], negative_qty_rows=excel_rows['negative_rows'],
    negative_rows_net_effect=r2(sum(n['net'] for n in neg_rows)),
    issue_rows=dict((k, sorted(er for i in v for er in row_of[i][:1])) for k, v in issues.items()
                    if k not in ('exact_duplicate_row', 'negative_qty_return_row')),
    issue_counts=dict((k, len(v)) for k, v in issues.items()),
    duplicate_source_rows=sorted(row_of[i][0] for i in issues['exact_duplicate_row']),
    bpk308_lines=len(bpk308_idx), bpk308_rows=sorted(row_of[i][0] for i in bpk308_idx),
    bpk308_completed_net=r2(sum(truth[i]['net'] for i in bpk308_idx if truth[i]['Status'] == 'Completed')),
    unmatched_sku_before_trim=len(bpk308_idx) + 12, unmatched_sku_after_trim=len(bpk308_idx),
    text_dates=55, sum_quantity_column_as_excel_sees_it=qty_numeric_sum,
    true_units_all_lines=sum(l['Quantity'] for l in truth),
    status_counts_truth=dict(Counter(l['Status'] for l in truth)),
    b2_region_line_counts=dict(sorted(Counter(pop_b2).items())),
    completed_lines=len(completed),
    completed_net_total=r2(sum(l['net'] for l in completed)),
    cancelled_net_total=r2(sum(l['net'] for l in truth if l['Status'] == 'Cancelled')),
    region_month_net=dict(('%s|%s' % (k[0], k[1].isoformat()[:7]), r2(v)) for k, v in sorted(reg_month.items())),
    region_q2_net=dict((reg, r2(sum(v for k, v in reg_month.items() if k[0] == reg)))
                       for reg in REGIONS + ['(blank)']),
    targets=dict(('%s|%s' % (k[0], k[1].isoformat()[:7]), v) for k, v in sorted(targets.items())),
    target_q2=dict((reg, sum(v for k, v in targets.items() if k[0] == reg)) for reg in REGIONS),
    attainment_q2_pct=dict((reg, round(100 * sum(v for k, v in reg_month.items() if k[0] == reg) /
                                       sum(v for k, v in targets.items() if k[0] == reg), 2)) for reg in REGIONS),
    attainment_by_month_pct=dict(('%s|%s' % (reg, m.isoformat()[:7]),
                                  round(100 * reg_month[(reg, m)] / targets[(reg, m)], 2))
                                 for reg in REGIONS for m in MONTHS_Q2),
    margin_by_category_excl_bpk308=margin,
    price_anomaly_rows=sorted(row_of[i][0] for i in issues['unit_price_x100_typo']),
    price_anomaly_overstatement=r2(typo_impact),
    top5_customers_correct=top5(cust_net),
    top5_customers_if_typos_uncorrected=top5(cust_net_typo),
    naive_sum_qty_x_price_x_1_minus_discount_all_rows=r2(naive),
)

# =========================================================================== POWER QUERY FILES (Part D)
PQ_DIR = os.path.join(ROOT, 'datasets', 'power_query')
EXP_DIR = os.path.join(PQ_DIR, 'cedarline_pos_exports')
CATS = ['Accessories', 'Apparel', 'Backpacks', 'Cooking', 'Footwear', 'Sleeping Gear', 'Tents']
agg = {}
for o in orders:
    if o['channel'] != 'In-Store' or o['status'] != 'completed' or not (date(2026, 6, 1) <= o['date'] <= DATA_END):
        continue
    for it in o['items']:
        c = PROD[it['product_id']]['category']
        k = (o['date'], o['store_id'], c)
        a = agg.setdefault(k, dict(orders=set(), units=0, gross=0.0, disc=0.0))
        a['orders'].add(o['order_id'])
        a['units'] += it['quantity']
        a['gross'] += it['quantity'] * it['unit_price']
        a['disc'] += it['quantity'] * it['unit_price'] * it['discount_pct']
pq_rows = []
for k in sorted(agg):
    a = agg[k]
    gross, disc = r2(a['gross']), r2(a['disc'])
    pq_rows.append(dict(date=k[0], store=k[1], cat=k[2], trans=len(a['orders']), units=a['units'],
                        gross=gross, disc=disc, net=r2(gross - disc)))


def write_csv(path, lines):
    with open(path, 'w', encoding='utf-8', newline='') as f:
        csv.writer(f, lineterminator='\r\n').writerows(lines)


jun = [r for r in pq_rows if r['date'].month == 6]
jul = [r for r in pq_rows if r['date'].month == 7]
aug = [r for r in pq_rows if r['date'].month == 8]
hdr = ['Date', 'StoreCode', 'Category', 'Transactions', 'Units', 'GrossSales', 'DiscountAmount', 'NetSales']
write_csv(os.path.join(EXP_DIR, 'POS_Summary_2026-06.csv'),
          [hdr] + [[r['date'].isoformat(), r['store'], r['cat'], r['trans'], r['units'],
                    '%.2f' % r['gross'], '%.2f' % r['disc'], '%.2f' % r['net']] for r in jun])
write_csv(os.path.join(EXP_DIR, 'POS_Summary_2026-07.csv'),
          [['Cedarline POS - Monthly Summary Export'], ['Generated 2026-08-01 06:00 by POS v4.2'], []] +
          [['Date', 'StoreCode', 'Category', 'Transactions', 'Qty Sold', 'GrossSales', 'DiscountAmount', 'NetSales']] +
          [[r['date'].strftime('%d/%m/%Y'), r['store'], r['cat'], r['trans'], r['units'],
            '%.2f' % r['gross'], '%.2f' % r['disc'], '%.2f' % r['net']] for r in jul])
aug_lines = [hdr + ['Notes']]
for r in aug:
    store = r['store'].lower() if (r['store'] == 'W02' and r['date'].day >= 16) else r['store']
    note = 'Promo weekend' if r['date'].day in (14, 15, 16) else ''
    aug_lines.append([r['date'].isoformat(), store, r['cat'], r['trans'], r['units'],
                      '%.2f' % r['gross'], '%.2f' % r['disc'], '%.2f' % r['net'], note])
aug_lines.append([])
aug_lines.append(['TOTAL', '', '', sum(r['trans'] for r in aug), sum(r['units'] for r in aug),
                  '%.2f' % sum(r['gross'] for r in aug), '%.2f' % sum(r['disc'] for r in aug),
                  '%.2f' % sum(r['net'] for r in aug), ''])
write_csv(os.path.join(EXP_DIR, 'POS_Summary_2026-08.csv'), aug_lines)
write_csv(os.path.join(PQ_DIR, 'stores.csv'),
          [['StoreCode', 'StoreName', 'Region', 'OpenedOn']] +
          [[s[0], s[1], s[2], s[3].isoformat()] for s in STORES])

store_month = defaultdict(lambda: dict(net=0.0, units=0))
for r in pq_rows:
    k = (r['store'], r['date'].month)
    store_month[k]['net'] += r['net']
    store_month[k]['units'] += r['units']
STORE_FACTOR = {'N01': 0.96, 'S01': 1.03, 'E01': 0.94, 'W01': 1.05, 'W02': 1.30}
store_targets = {}
for s in STORES:
    for m in (6, 7, 8):
        store_targets[(s[0], m)] = round(store_month[(s[0], m)]['net'] * STORE_FACTOR[s[0]] *
                                         rng.uniform(0.98, 1.02) / 500.0) * 500
write_csv(os.path.join(PQ_DIR, 'store_targets_summer_2026.csv'),
          [['StoreCode', '2026-06', '2026-07', '2026-08']] +
          [[s[0]] + [store_targets[(s[0], m)] for m in (6, 7, 8)] for s in STORES])
pq_gt = dict(
    rows_per_file=dict(jun=len(jun), jul=len(jul), aug=len(aug)), total_data_rows=len(pq_rows),
    aug_rows_with_lowercase_w02=sum(1 for r in aug if r['store'] == 'W02' and r['date'].day >= 16),
    store_month=[dict(store=s[0], month='2026-%02d' % m, net_sales=r2(store_month[(s[0], m)]['net']),
                      units=store_month[(s[0], m)]['units'], target=store_targets[(s[0], m)],
                      variance=r2(store_month[(s[0], m)]['net'] - store_targets[(s[0], m)]),
                      attainment_pct=round(100 * store_month[(s[0], m)]['net'] / store_targets[(s[0], m)], 2))
                 for s in STORES for m in (6, 7, 8)],
    month_net=dict(('2026-%02d' % m, r2(sum(v['net'] for k, v in store_month.items() if k[1] == m)))
                   for m in (6, 7, 8)),
    july_rows_day_le_12=sum(1 for r in jul if r['date'].day <= 12),
    july_net_day_le_12=r2(sum(r['net'] for r in jul if r['date'].day <= 12)),
)

# =========================================================================== DB summary + ground truth file
db_gt = dict(
    customers=len(records), persons=len(persons), duplicate_records=len(records) - len(persons),
    orders=len(orders), order_items=len(items), returns=len(returns),
    walk_in_orders=sum(1 for o in orders if o['customer_id'] is None),
    header_only_orders=sum(1 for o in orders if o.get('header_only')),
    status_counts=dict(Counter(o['status'] for o in orders)),
    orders_per_year=dict(Counter(str(o['date'].year) for o in orders)),
)
gt = dict(seed=SEED, db=db_gt, excel=excel_gt, power_query=pq_gt)
out = os.path.join(ROOT, 'solutions', '00_diagnostic', '_generated', 'ground_truth.json')
with open(out, 'w', encoding='utf-8') as f:
    json.dump(gt, f, indent=2, default=str)
print(json.dumps(dict(db=db_gt, excel_lines=len(truth), excel_last_row=last_row,
                      pq_rows=pq_gt['rows_per_file']), indent=2))
